"""Excel 보고서 자동 생성 - KHIDI 양식 기반

실제 보안 컨설팅 현장에서 납품하는 양식 그대로 자동 생성합니다.
구조: 표지 → 점검대상 → 요약(매트릭스) → 그래프 → 장비별 상세 시트
"""
import io
from datetime import datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# === 스타일 정의 ===
FONT_TITLE = Font(name="맑은 고딕", size=22, bold=True)
FONT_HEADER = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
FONT_BODY = Font(name="맑은 고딕", size=10)
FONT_SMALL = Font(name="맑은 고딕", size=9, color="666666")
FONT_PASS = Font(name="맑은 고딕", size=10, color="2E7D32", bold=True)
FONT_FAIL = Font(name="맑은 고딕", size=10, color="C62828", bold=True)

FILL_HEADER = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
FILL_PASS = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
FILL_FAIL = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
FILL_NA = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
FILL_LIGHT = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
FILL_CATEGORY = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")

BORDER_THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


class ExcelReportGenerator:
    """KHIDI 양식 기반 Excel 취약점 진단 보고서 생성기"""

    def generate(self, report_data: dict) -> bytes:
        """보고서 생성 → bytes 반환
        
        report_data 구조:
        {
            "project_name": str,
            "client_name": str,
            "consultant": str,
            "report_date": str,
            "assets": [
                {
                    "name": str, "ip": str, "vendor": str,
                    "hostname": str, "asset_type": str, "manager": str,
                    "results": [
                        {"category": str, "item_id": str, "title": str,
                         "risk_level": str, "verdict": str, "detail": str}
                    ]
                }
            ]
        }
        """
        wb = Workbook()
        wb.remove(wb.active)  # 기본 시트 제거

        assets = report_data.get("assets", [])

        self._create_cover(wb, report_data)
        self._create_target_list(wb, assets)
        self._create_summary_matrix(wb, assets)
        self._create_chart_sheet(wb, assets)

        for asset in assets:
            self._create_asset_sheet(wb, asset)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()

    def _create_cover(self, wb: Workbook, data: dict):
        """표지 시트"""
        ws = wb.create_sheet("표지")
        ws.sheet_properties.tabColor = "2C3E50"

        ws.merge_cells("B4:H4")
        ws["B4"] = "기술적 취약점 진단 결과 보고서"
        ws["B4"].font = FONT_TITLE
        ws["B4"].alignment = Alignment(horizontal="center")

        info = [
            ("프로젝트", data.get("project_name", "")),
            ("고객사", data.get("client_name", "")),
            ("진단일", data.get("report_date", datetime.now().strftime("%Y-%m-%d"))),
            ("수행사", data.get("consultant", "로뎀시스템")),
            ("진단 기준", "주요정보통신기반시설 기술적 취약점 분석·평가 방법 상세가이드"),
            ("총 자산 수", str(len(data.get("assets", [])))),
        ]
        for i, (label, value) in enumerate(info):
            row = 8 + i
            ws.merge_cells(f"C{row}:D{row}")
            ws[f"C{row}"] = label
            ws[f"C{row}"].font = Font(name="맑은 고딕", size=11, bold=True)
            ws[f"C{row}"].border = BORDER_THIN
            ws.merge_cells(f"E{row}:G{row}")
            ws[f"E{row}"] = value
            ws[f"E{row}"].font = Font(name="맑은 고딕", size=11)
            ws[f"E{row}"].border = BORDER_THIN

        ws.column_dimensions["B"].width = 4
        for col in "CDEFGH":
            ws.column_dimensions[col].width = 15

    def _create_target_list(self, wb: Workbook, assets: list):
        """점검대상 시트"""
        ws = wb.create_sheet("점검대상")
        ws.sheet_properties.tabColor = "3498DB"

        ws.merge_cells("A1:H1")
        ws["A1"] = "점검 대상"
        ws["A1"].font = Font(name="맑은 고딕", size=14, bold=True)

        headers = ["No", "구분", "자산명(용도)", "IP", "Vendor", "호스트명", "담당자", "비고"]
        widths = [6, 14, 25, 18, 15, 15, 10, 15]
        for col_idx, (h, w) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(row=3, column=col_idx, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN
            ws.column_dimensions[get_column_letter(col_idx)].width = w

        for i, asset in enumerate(assets):
            row = 4 + i
            vals = [
                i + 1,
                self._type_label(asset.get("asset_type", "")),
                asset.get("name", ""),
                asset.get("ip", ""),
                asset.get("vendor", ""),
                asset.get("hostname", ""),
                asset.get("manager", ""),
                "",
            ]
            for col_idx, val in enumerate(vals, 1):
                cell = ws.cell(row=row, column=col_idx, value=val)
                cell.font = FONT_BODY
                cell.border = BORDER_THIN
                cell.alignment = ALIGN_CENTER if col_idx in (1, 2) else ALIGN_LEFT

    def _create_summary_matrix(self, wb: Workbook, assets: list):
        """요약 매트릭스 - 전체 항목 x 전체 장비 교차표"""
        ws = wb.create_sheet("요약")
        ws.sheet_properties.tabColor = "E74C3C"

        ws.merge_cells("A1:F1")
        ws["A1"] = "취약점 진단 결과 요약"
        ws["A1"].font = Font(name="맑은 고딕", size=14, bold=True)

        # 헤더: 구분, 카테고리, 항목코드, 점검항목, 중요도, 지수, [장비별...]
        base_headers = ["구분", "카테고리", "항목코드", "점검항목", "중요도", "지수"]
        asset_names = [a.get("name", f"자산{i+1}") for i, a in enumerate(assets)]
        all_headers = base_headers + asset_names

        for col_idx, h in enumerate(all_headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN

        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 30
        ws.column_dimensions["E"].width = 8
        ws.column_dimensions["F"].width = 8
        for i in range(len(asset_names)):
            ws.column_dimensions[get_column_letter(7 + i)].width = 12

        # 전체 항목 수집
        all_items = {}
        for asset in assets:
            for r in asset.get("results", []):
                item_id = r.get("item_id", "")
                if item_id not in all_items:
                    all_items[item_id] = {
                        "category": r.get("category", ""),
                        "title": r.get("title", ""),
                        "risk_level": r.get("risk_level", ""),
                    }

        row = 4
        risk_score = {"상": 10, "중": 5, "하": 2, "high": 10, "medium": 5, "low": 2}
        prev_cat = ""

        for item_id in sorted(all_items.keys()):
            info = all_items[item_id]
            cat = info["category"]
            cat_display = cat if cat != prev_cat else ""
            prev_cat = cat

            ws.cell(row=row, column=1, value=cat_display).font = FONT_BODY
            ws.cell(row=row, column=2, value=cat).font = FONT_BODY
            ws.cell(row=row, column=3, value=item_id).font = Font(name="맑은 고딕", size=10, bold=True)
            ws.cell(row=row, column=4, value=info["title"]).font = FONT_BODY
            ws.cell(row=row, column=5, value=info["risk_level"]).font = FONT_BODY
            ws.cell(row=row, column=6, value=risk_score.get(info["risk_level"], 5)).font = FONT_BODY

            for asset_idx, asset in enumerate(assets):
                verdict = ""
                for r in asset.get("results", []):
                    if r.get("item_id") == item_id:
                        verdict = "양호" if r["verdict"] == "pass" else "취약" if r["verdict"] == "fail" else "N/A"
                        break

                cell = ws.cell(row=row, column=7 + asset_idx, value=verdict or "-")
                if verdict == "양호":
                    cell.font = FONT_PASS
                    cell.fill = FILL_PASS
                elif verdict == "취약":
                    cell.font = FONT_FAIL
                    cell.fill = FILL_FAIL
                else:
                    cell.font = FONT_SMALL
                    cell.fill = FILL_NA
                cell.alignment = ALIGN_CENTER
                cell.border = BORDER_THIN

            for col_idx in range(1, 7):
                ws.cell(row=row, column=col_idx).border = BORDER_THIN
                ws.cell(row=row, column=col_idx).alignment = ALIGN_CENTER if col_idx in (1, 3, 5, 6) else ALIGN_LEFT
                if cat_display:
                    ws.cell(row=row, column=1).fill = FILL_CATEGORY

            row += 1

    def _create_chart_sheet(self, wb: Workbook, assets: list):
        """그래프 시트 - 장비별 준수율"""
        ws = wb.create_sheet("그래프")
        ws.sheet_properties.tabColor = "F39C12"

        ws["B1"] = "장비별 준수율 그래프"
        ws["B1"].font = Font(name="맑은 고딕", size=14, bold=True)

        headers = ["No.", "장비명", "준수율"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.border = BORDER_THIN

        for i, asset in enumerate(assets):
            results = asset.get("results", [])
            total = len(results) or 1
            passed = sum(1 for r in results if r.get("verdict") == "pass")
            rate = round(passed / total, 4)

            ws.cell(row=4 + i, column=1, value=i + 1).border = BORDER_THIN
            ws.cell(row=4 + i, column=2, value=asset.get("name", "")).border = BORDER_THIN
            cell = ws.cell(row=4 + i, column=3, value=rate)
            cell.number_format = "0.0%"
            cell.border = BORDER_THIN

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 12

        if assets:
            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = "장비별 취약점 준수율"
            chart.y_axis.title = "준수율"
            chart.y_axis.scaling.min = 0
            chart.y_axis.scaling.max = 1
            chart.y_axis.numFmt = "0%"

            data = Reference(ws, min_col=3, min_row=3, max_row=3 + len(assets))
            cats = Reference(ws, min_col=2, min_row=4, max_row=3 + len(assets))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.width = 25
            chart.height = 15

            ws.add_chart(chart, "B" + str(6 + len(assets)))

    def _create_asset_sheet(self, wb: Workbook, asset: dict):
        """개별 장비 시트 - KHIDI 양식 6컬럼"""
        name = asset.get("name", "자산")[:31]  # 시트명 31자 제한
        ws = wb.create_sheet(name)

        # 헤더 정보
        info = [
            ("자산명", asset.get("name", "")),
            ("IP", asset.get("ip", "")),
            ("Vendor", asset.get("vendor", "")),
        ]
        for i, (label, value) in enumerate(info):
            ws.cell(row=1 + i, column=1, value=label).font = Font(name="맑은 고딕", size=10, bold=True)
            ws.merge_cells(start_row=1 + i, start_column=3, end_row=1 + i, end_column=6)
            ws.cell(row=1 + i, column=3, value=value).font = FONT_BODY

        # 컬럼 헤더 (Row 5)
        headers = ["구분", "코드", "점 검 항 목", "위험도", "점검\n결과", "현황"]
        widths = [10, 8, 35, 8, 8, 50]
        for col_idx, (h, w) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN
            ws.column_dimensions[get_column_letter(col_idx)].width = w

        # 데이터
        prev_cat = ""
        row = 6
        for r in asset.get("results", []):
            cat = r.get("category", "")
            cat_display = cat if cat != prev_cat else ""
            prev_cat = cat

            verdict = "양호" if r["verdict"] == "pass" else "취약" if r["verdict"] == "fail" else "N/A"
            detail = r.get("detail", "")

            values = [cat_display, r.get("item_id", ""), r.get("title", ""), r.get("risk_level", ""), verdict, detail]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=val)
                cell.font = FONT_BODY
                cell.border = BORDER_THIN
                cell.alignment = ALIGN_CENTER if col_idx in (1, 2, 4, 5) else ALIGN_LEFT

                if col_idx == 5:
                    if verdict == "양호":
                        cell.font = FONT_PASS
                        cell.fill = FILL_PASS
                    elif verdict == "취약":
                        cell.font = FONT_FAIL
                        cell.fill = FILL_FAIL

                if col_idx == 1 and cat_display:
                    cell.fill = FILL_CATEGORY
                    cell.font = Font(name="맑은 고딕", size=10, bold=True)

            row += 1

        # 하단 통계
        results = asset.get("results", [])
        total = len(results)
        passed = sum(1 for r in results if r.get("verdict") == "pass")
        failed = sum(1 for r in results if r.get("verdict") == "fail")

        row += 1
        ws.cell(row=row, column=1, value="합계").font = Font(name="맑은 고딕", size=10, bold=True)
        ws.cell(row=row, column=3, value=f"전체 {total}개 항목").font = FONT_BODY
        ws.cell(row=row, column=4, value=f"양호 {passed}").font = FONT_PASS
        ws.cell(row=row, column=5, value=f"취약 {failed}").font = FONT_FAIL
        rate = round(100 * passed / total, 1) if total else 0
        ws.cell(row=row, column=6, value=f"준수율 {rate}%").font = Font(name="맑은 고딕", size=10, bold=True)

    def _type_label(self, asset_type: str) -> str:
        labels = {
            "server_unix": "서버(Unix/Linux)",
            "server_windows": "서버(Windows)",
            "dbms": "DBMS",
            "network": "네트워크장비",
            "security": "보안장비",
            "web": "웹서비스",
            "pc": "PC",
            "cloud": "클라우드",
        }
        return labels.get(asset_type, asset_type)
